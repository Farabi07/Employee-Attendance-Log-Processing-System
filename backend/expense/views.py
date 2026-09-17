from datetime import date, timedelta
from decimal import Decimal
import os
import re

from django.db import transaction
from django.db.models import Sum
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.decorators import api_view, parser_classes, permission_classes
from rest_framework.parsers import MultiPartParser, JSONParser
from rest_framework.response import Response

from authentication.models import Employee
from authentication.permissions import HasActiveSubscription, IsManagerOrModerator
from attendance.models import Attendance
from .models import Expense
from .serializers import ExpenseSerializer


def _period_bounds(period):
    today = timezone.localdate()
    if period == "daily":
        return today, today
    if period == "yearly":
        return date(today.year, 1, 1), date(today.year, 12, 31)
    return date(today.year, today.month, 1), date(today.year + (today.month == 12), (today.month % 12) + 1, 1) - timedelta(days=1)


def _base_queryset(request):
    return Expense.objects.filter(organization=request.user.organization).select_related("recipient", "branch")


def _revenue(organization, start, end):
    return Attendance.objects.filter(employee__organization=organization, date__range=(start, end)).aggregate(total=Sum("earnings"))["total"] or Decimal("0")


@api_view(["GET"])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def expense_summary(request):
    start, end = _period_bounds(request.query_params.get("period", "monthly"))
    qs = _base_queryset(request).filter(date__range=(start, end))
    total = qs.aggregate(total=Sum("amount"))["total"] or Decimal("0")
    category_totals = qs.values("category").annotate(total=Sum("amount")).order_by("-total")
    revenue = _revenue(request.user.organization, start, end)
    return Response({
        "period": request.query_params.get("period", "monthly"),
        "date_from": start,
        "date_to": end,
        "revenue": revenue,
        "total_expense": total,
        "profit": revenue - total,
        "categories": list(category_totals),
    })


@api_view(["GET"])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def expense_list(request):
    start, end = _period_bounds(request.query_params.get("period", "monthly"))
    expenses = _base_queryset(request).filter(date__range=(start, end))
    return Response({"expenses": ExpenseSerializer(expenses[:500], many=True).data})


@api_view(["POST"])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def expense_create(request):
    serializer = ExpenseSerializer(data=request.data, context={"request": request})
    serializer.is_valid(raise_exception=True)
    serializer.save(
        organization=request.user.organization,
        branch=getattr(request.user, "branch", None),
        created_by=request.user,
        source="direct",
    )
    return Response(serializer.data, status=status.HTTP_201_CREATED)


@api_view(["POST"])
@parser_classes([MultiPartParser])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def expense_import_excel(request):
    upload = request.FILES.get("file")
    if not upload:
        return Response({"detail": "Attach an Excel file in the file field."}, status=status.HTTP_400_BAD_REQUEST)
    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return Response({"detail": "The spreadsheet is empty."}, status=status.HTTP_400_BAD_REQUEST)
        headers = {str(value).strip().lower(): index for index, value in enumerate(rows[0]) if value is not None}
        required = {"amount", "category"}
        missing = required - headers.keys()
        if missing:
            return Response({"detail": f"Missing columns: {', '.join(sorted(missing))}"}, status=status.HTTP_400_BAD_REQUEST)
        created = []
        with transaction.atomic():
            for line_number, row in enumerate(rows[1:], 2):
                if not any(value is not None for value in row):
                    continue
                amount = row[headers["amount"]]
                category = str(row[headers["category"]]).strip().lower()
                if not amount or category not in Expense.Category.values:
                    raise ValueError(f"Row {line_number}: amount and a valid category are required.")
                recipient_id = row[headers["recipient_id"]] if "recipient_id" in headers else None
                if recipient_id and not Employee.objects.filter(pk=recipient_id, organization=request.user.organization).exists():
                    raise ValueError(f"Row {line_number}: recipient is not in this store.")
                created.append(Expense(
                    organization=request.user.organization,
                    branch=getattr(request.user, "branch", None),
                    recipient_id=recipient_id,
                    amount=amount,
                    category=category,
                    description=str(row[headers["description"]]).strip() if "description" in headers and row[headers["description"]] else "",
                    date=row[headers["date"]] if "date" in headers and row[headers["date"]] else timezone.localdate(),
                    source="excel",
                    created_by=request.user,
                ))
            Expense.objects.bulk_create(created)
        return Response({"created": len(created)}, status=status.HTTP_201_CREATED)
    except (ValueError, TypeError, KeyError) as exc:
        return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)


@api_view(["GET"])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def expense_template(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["date", "amount", "category", "description", "recipient_id"])
    sheet.append([timezone.localdate(), 0, "supplies", "Example", ""])
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="expense-template.xlsx"'
    workbook.save(response)
    return response


def _receipt_amount(lines):
    candidates = []
    for line in lines:
        match = re.search(r"(?:total|grand total|amount due|net total)\D{0,12}(\d+[,.]?\d{0,2})", line, re.I)
        if match:
            candidates.append(match.group(1).replace(",", ""))
    if candidates:
        return candidates[-1]
    values = re.findall(r"(?<!\d)(\d{1,7}[.]\d{2})(?!\d)", " ".join(lines))
    return values[-1] if values else ""


def _receipt_date(lines):
    for line in lines:
        match = re.search(r"(\d{4})[-/]([01]?\d)[-/]([0-3]?\d)|(\d{1,2})[-/]([01]?\d)[-/](\d{2,4})", line)
        if not match:
            continue
        groups = match.groups()
        if groups[0]:
            return f"{int(groups[0]):04d}-{int(groups[1]):02d}-{int(groups[2]):02d}"
        year = int(groups[5])
        year += 2000 if year < 100 else 0
        return f"{year:04d}-{int(groups[4]):02d}-{int(groups[3]):02d}"
    return str(timezone.localdate())


def _receipt_category(text):
    lowered = text.lower()
    for keyword, category in (
        ("rent", "rent"), ("utility", "utilities"), ("electric", "utilities"),
        ("transport", "transport"), ("fuel", "transport"), ("market", "supplies"),
        ("supply", "supplies"), ("salary", "salary"), ("wage", "salary"),
        ("advert", "marketing"),
    ):
        if keyword in lowered:
            return category
    return "other"


@api_view(["POST"])
@parser_classes([MultiPartParser])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def expense_receipt_extract(request):
    upload = request.FILES.get("file")
    if not upload:
        return Response({"detail": "Attach a receipt image in the file field."}, status=status.HTTP_400_BAD_REQUEST)
    if not os.environ.get("AWS_ACCESS_KEY_ID") or not os.environ.get("AWS_SECRET_ACCESS_KEY"):
        return Response({"detail": "Receipt OCR is not configured. Set AWS Textract credentials on the backend."}, status=status.HTTP_503_SERVICE_UNAVAILABLE)
    try:
        import boto3

        client = boto3.client("textract", region_name=os.environ.get("AWS_REGION", "us-east-1"))
        result = client.detect_document_text(Document={"Bytes": upload.read()})
        lines = [block["Text"] for block in result.get("Blocks", []) if block.get("BlockType") == "LINE"]
        if not lines:
            return Response({"detail": "No text was detected in the receipt."}, status=status.HTTP_422_UNPROCESSABLE_ENTITY)
        text = " ".join(lines)
        return Response({
            "amount": _receipt_amount(lines),
            "date": _receipt_date(lines),
            "category": _receipt_category(text),
            "description": next((line[:255] for line in lines if len(line) > 3), "Receipt expense"),
            "text": "\n".join(lines),
        })
    except Exception as exc:
        return Response({"detail": f"Receipt OCR failed: {exc}"}, status=status.HTTP_502_BAD_GATEWAY)
