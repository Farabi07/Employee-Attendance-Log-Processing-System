from datetime import date, timedelta
from decimal import Decimal
import csv
import io
import json
import os
import re

from django.db import transaction
from django.db.models import Sum
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework import status
from rest_framework.decorators import api_view, parser_classes, permission_classes
from rest_framework.parsers import MultiPartParser, JSONParser
from rest_framework.response import Response

from authentication.models import Employee
from authentication.permissions import HasActiveSubscription, IsManagerOrModerator
from .models import Expense, ExpenseCategory, FinanceAuditLog, Income
from .serializers import ExpenseCategorySerializer, ExpenseSerializer, IncomeSerializer
from ..filters import ExpenseFilter


def _period_bounds(period):
    today = timezone.localdate()
    if period == "daily":
        return today, today
    if period == "yearly":
        return date(today.year, 1, 1), date(today.year, 12, 31)
    return date(today.year, today.month, 1), date(today.year + (today.month == 12), (today.month % 12) + 1, 1) - timedelta(days=1)


def _base_queryset(request):
    return Expense.objects.filter(organization=request.user.organization).select_related("recipient", "branch")


def _audit(request, entity_type, entity_id, action, changes=None):
    FinanceAuditLog.objects.create(
        organization=request.user.organization,
        actor=request.user,
        entity_type=entity_type,
        entity_id=entity_id,
        action=action,
        changes=json.loads(json.dumps(changes or {}, default=str)),
    )


def _ensure_category(organization, name):
    """Reuse a category case-insensitively, creating it when an import adds a new one."""
    normalized = str(name).strip()
    existing = ExpenseCategory.objects.filter(organization=organization, name__iexact=normalized).first()
    if existing:
        return existing
    return ExpenseCategory.objects.create(organization=organization, name=normalized)


def _filtered_expenses(request):
    qs = _base_queryset(request)
    qs = ExpenseFilter(request.query_params, queryset=qs).qs
    period = request.query_params.get("period")
    if period:
        start, end = _period_bounds(period)
        qs = qs.filter(date__range=(start, end))
    if request.query_params.get("date_from"):
        qs = qs.filter(date__gte=request.query_params["date_from"])
    if request.query_params.get("date_to"):
        qs = qs.filter(date__lte=request.query_params["date_to"])
    if request.query_params.get("category"):
        qs = qs.filter(category=request.query_params["category"])
    return qs


def _revenue(organization, start, end):
    return Income.objects.filter(organization=organization, date__range=(start, end)).aggregate(total=Sum("amount"))["total"] or Decimal("0")


@api_view(["GET"])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def expense_summary(request):
    start, end = _period_bounds(request.query_params.get("period", "monthly"))
    qs = _filtered_expenses(request).filter(date__range=(start, end))
    total = qs.aggregate(total=Sum("amount"))["total"] or Decimal("0")
    category_totals = qs.exclude(category="profit").values("category").annotate(total=Sum("amount")).order_by("-total")
    sales_category_totals = Income.objects.filter(
        organization=request.user.organization,
        date__range=(start, end),
    ).values("category").annotate(total=Sum("amount")).order_by("-total")
    revenue = _revenue(request.user.organization, start, end)
    return Response({
        "period": request.query_params.get("period", "monthly"),
        "date_from": start,
        "date_to": end,
        "revenue": revenue,
        "total_expense": total,
        "profit": revenue - total,
        "categories": list(category_totals),
        "sales_categories": list(sales_category_totals),
    })


@api_view(["GET"])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def expense_list(request):
    expenses = _filtered_expenses(request)
    total = expenses.count()
    page = request.query_params.get("page")
    size = request.query_params.get("size")
    if page or size:
        from commons.pagination import Pagination
        pagination = Pagination()
        pagination.page = page
        pagination.size = size
        expenses = pagination.paginate_data(expenses)
        return Response({
            "expenses": ExpenseSerializer(expenses, many=True).data,
            "page": pagination.page,
            "size": pagination.size,
            "total_pages": pagination.total_pages,
            "total_elements": total,
        })
    return Response({"expenses": ExpenseSerializer(expenses[:500], many=True).data, "total_elements": total})


@api_view(["POST"])
@parser_classes([MultiPartParser, JSONParser])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def expense_create(request):
    serializer = ExpenseSerializer(data=request.data, context={"request": request})
    serializer.is_valid(raise_exception=True)
    serializer.save(
        organization=request.user.organization,
        branch=getattr(request.user, "branch", None),
        created_by=request.user,
        source="direct",
    )
    _audit(request, "expense", serializer.instance.pk, "create", serializer.validated_data)
    return Response(serializer.data, status=status.HTTP_201_CREATED)


@api_view(["GET", "PATCH", "DELETE"])
@parser_classes([MultiPartParser, JSONParser])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def expense_detail(request, pk):
    expense = _base_queryset(request).filter(pk=pk).first()
    if not expense:
        return Response({"detail": "Expense not found."}, status=status.HTTP_404_NOT_FOUND)
    if request.method == "GET":
        return Response(ExpenseSerializer(expense, context={"request": request}).data)
    if request.method == "DELETE":
        _audit(request, "expense", expense.pk, "delete")
        expense.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)
    serializer = ExpenseSerializer(expense, data=request.data, partial=True, context={"request": request})
    serializer.is_valid(raise_exception=True)
    serializer.save()
    _audit(request, "expense", expense.pk, "update", serializer.validated_data)
    return Response(serializer.data)


@api_view(["GET", "POST"])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def expense_categories(request):
    qs = ExpenseCategory.objects.filter(organization=request.user.organization)
    if request.method == "GET":
        return Response(ExpenseCategorySerializer(qs, many=True).data)
    serializer = ExpenseCategorySerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    serializer.save(organization=request.user.organization)
    _audit(request, "category", serializer.instance.pk, "create", serializer.validated_data)
    return Response(serializer.data, status=status.HTTP_201_CREATED)


@api_view(["PATCH", "DELETE"])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def expense_category_detail(request, pk):
    category = ExpenseCategory.objects.filter(pk=pk, organization=request.user.organization).first()
    if not category:
        return Response({"detail": "Category not found."}, status=status.HTTP_404_NOT_FOUND)
    if request.method == "DELETE":
        category.is_active = False
        category.save(update_fields=("is_active",))
        return Response(status=status.HTTP_204_NO_CONTENT)
    serializer = ExpenseCategorySerializer(category, data=request.data, partial=True)
    serializer.is_valid(raise_exception=True)
    serializer.save()
    return Response(serializer.data)


@api_view(["POST"])
@parser_classes([MultiPartParser])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def expense_import_excel(request):
    upload = request.FILES.get("file")
    if not upload:
        return Response({"detail": "Attach an Excel file in the file field."}, status=status.HTTP_400_BAD_REQUEST)
    try:
        if upload.name.lower().endswith(".csv"):
            rows = list(csv.reader(io.TextIOWrapper(upload.file, encoding="utf-8-sig")))
        else:
            workbook = load_workbook(upload, read_only=True, data_only=True)
            rows = list(workbook.active.iter_rows(values_only=True))
        if not rows:
            return Response({"detail": "The spreadsheet is empty."}, status=status.HTTP_400_BAD_REQUEST)
        headers = {str(value).strip().lower(): index for index, value in enumerate(rows[0]) if value is not None}
        required = {"amount", "category"}
        missing = required - headers.keys()
        if missing:
            return Response({"detail": f"Missing columns: {', '.join(sorted(missing))}"}, status=status.HTTP_400_BAD_REQUEST)
        created = []
        with transaction.atomic():
            for line_number, row in enumerate(rows[1:], 2):
                if not any(value is not None for value in row):
                    continue
                amount = row[headers["amount"]]
                category = str(row[headers["category"]]).strip().lower()
                if not amount or not category:
                    raise ValueError(f"Row {line_number}: amount and category are required.")
                if category == "profit":
                    raise ValueError(f"Row {line_number}: profit cannot be used as an expense category.")
                category_ref = _ensure_category(request.user.organization, category)
                recipient_id = row[headers["recipient_id"]] if "recipient_id" in headers else None
                if recipient_id and not Employee.objects.filter(pk=recipient_id, organization=request.user.organization).exists():
                    raise ValueError(f"Row {line_number}: recipient is not in this store.")
                created.append(Expense(
                    organization=request.user.organization,
                    branch=getattr(request.user, "branch", None),
                    recipient_id=recipient_id,
                    amount=amount,
                    category=category,
                    category_ref=category_ref,
                    description=str(row[headers["description"]]).strip() if "description" in headers and row[headers["description"]] else "",
                    date=row[headers["date"]] if "date" in headers and row[headers["date"]] else timezone.localdate(),
                    source="excel",
                    created_by=request.user,
                ))
            Expense.objects.bulk_create(created)
        return Response({"created": len(created)}, status=status.HTTP_201_CREATED)
    except (ValueError, TypeError, KeyError) as exc:
        return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)


@api_view(["POST"])
@parser_classes([MultiPartParser])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def finance_workbook_import(request):
    """Import one workbook containing Sales and Expenses sheets."""
    upload = request.FILES.get("file")
    if not upload:
        return Response({"detail": "Attach an Excel workbook in the file field."}, status=status.HTTP_400_BAD_REQUEST)

    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
        sheet_names = {name.lower(): name for name in workbook.sheetnames}
        if "sales" not in sheet_names or "expenses" not in sheet_names:
            return Response(
                {"detail": "Workbook must contain Sales and Expenses sheets."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        def read_sheet(name, required):
            rows = list(workbook[sheet_names[name]].iter_rows(values_only=True))
            if not rows:
                return []
            headers = {str(value).strip().lower(): index for index, value in enumerate(rows[0]) if value is not None}
            missing = required - headers.keys()
            if missing:
                raise ValueError(f"{name.title()} sheet is missing: {', '.join(sorted(missing))}")
            return headers, rows[1:]

        expense_headers, expense_rows = read_sheet("expenses", {"amount", "category"})
        sales_headers, sales_rows = read_sheet("sales", {"amount", "category"})
        expenses = []
        incomes = []

        with transaction.atomic():
            for line_number, row in enumerate(expense_rows, 2):
                if not any(value is not None for value in row):
                    continue
                amount = row[expense_headers["amount"]]
                category = str(row[expense_headers["category"]]).strip().lower()
                if not amount or not category:
                    raise ValueError(f"Expenses sheet row {line_number}: amount and category are required.")
                if category == "profit":
                    raise ValueError(f"Expenses sheet row {line_number}: profit cannot be used as an expense category.")
                category_ref = _ensure_category(request.user.organization, category)
                recipient_id = row[expense_headers["recipient_id"]] if "recipient_id" in expense_headers else None
                if recipient_id and not Employee.objects.filter(pk=recipient_id, organization=request.user.organization).exists():
                    raise ValueError(f"Expenses sheet row {line_number}: recipient is not in this store.")
                expenses.append(Expense(
                    organization=request.user.organization,
                    branch=getattr(request.user, "branch", None),
                    recipient_id=recipient_id,
                    amount=amount,
                    category=category,
                    category_ref=category_ref,
                    description=str(row[expense_headers["description"]]).strip() if "description" in expense_headers and row[expense_headers["description"]] else "",
                    vendor_name=str(row[expense_headers["vendor_name"]]).strip() if "vendor_name" in expense_headers and row[expense_headers["vendor_name"]] else "",
                    payment_method=str(row[expense_headers["payment_method"]]).strip() if "payment_method" in expense_headers and row[expense_headers["payment_method"]] else "cash",
                    date=row[expense_headers["date"]] if "date" in expense_headers and row[expense_headers["date"]] else timezone.localdate(),
                    source="excel",
                    created_by=request.user,
                ))

            for line_number, row in enumerate(sales_rows, 2):
                if not any(value is not None for value in row):
                    continue
                amount = row[sales_headers["amount"]]
                category = str(row[sales_headers["category"]]).strip().lower()
                if not amount or not category:
                    raise ValueError(f"Sales sheet row {line_number}: amount and category are required.")
                _ensure_category(request.user.organization, category)
                incomes.append(Income(
                    organization=request.user.organization,
                    branch=getattr(request.user, "branch", None),
                    amount=amount,
                    category=category,
                    description=str(row[sales_headers["description"]]).strip() if "description" in sales_headers and row[sales_headers["description"]] else "",
                    date=row[sales_headers["date"]] if "date" in sales_headers and row[sales_headers["date"]] else timezone.localdate(),
                    source="excel",
                    created_by=request.user,
                ))

            Expense.objects.bulk_create(expenses)
            Income.objects.bulk_create(incomes)
        return Response({"expenses_created": len(expenses), "sales_created": len(incomes)}, status=status.HTTP_201_CREATED)
    except (ValueError, TypeError, KeyError) as exc:
        return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)


@api_view(["GET"])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def expense_template(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["date", "amount", "category", "description", "recipient_id"])
    sheet.append([timezone.localdate(), 0, "supplies", "Example", ""])
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = 'attachment; filename="expense-template.xlsx"'
    workbook.save(response)
    return response


@api_view(["POST"])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def income_create(request):
    serializer = IncomeSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    serializer.save(organization=request.user.organization, branch=getattr(request.user, "branch", None), created_by=request.user, source="direct")
    _audit(request, "income", serializer.instance.pk, "create", serializer.validated_data)
    return Response(serializer.data, status=status.HTTP_201_CREATED)


@api_view(["GET"])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def income_list(request):
    start, end = _period_bounds(request.query_params.get("period", "monthly"))
    income = Income.objects.filter(organization=request.user.organization, date__range=(start, end))
    if request.query_params.get("date_from"):
        income = income.filter(date__gte=request.query_params["date_from"])
    if request.query_params.get("date_to"):
        income = income.filter(date__lte=request.query_params["date_to"])
    if request.query_params.get("category"):
        income = income.filter(category=request.query_params["category"])
    return Response({"income": IncomeSerializer(income[:500], many=True).data})


@api_view(["POST"])
@parser_classes([MultiPartParser])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def income_import_excel(request):
    upload = request.FILES.get("file")
    if not upload:
        return Response({"detail": "Attach an Excel file in the file field."}, status=status.HTTP_400_BAD_REQUEST)
    try:
        workbook = load_workbook(upload, read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))
        headers = {str(value).strip().lower(): index for index, value in enumerate(rows[0]) if value is not None} if rows else {}
        missing = {"amount", "category"} - headers.keys()
        if missing:
            return Response({"detail": f"Missing columns: {', '.join(sorted(missing))}"}, status=status.HTTP_400_BAD_REQUEST)
        created = []
        with transaction.atomic():
            for row in rows[1:]:
                if not any(value is not None for value in row):
                    continue
                category = str(row[headers["category"]]).strip().lower()
                if not category:
                    raise ValueError("Each sales row must include a category.")
                _ensure_category(request.user.organization, category)
                created.append(Income(organization=request.user.organization, branch=getattr(request.user, "branch", None), created_by=request.user, source="excel", amount=row[headers["amount"]], category=category, description=str(row[headers["description"]]).strip() if "description" in headers and row[headers["description"]] else "", date=row[headers["date"]] if "date" in headers and row[headers["date"]] else timezone.localdate()))
            Income.objects.bulk_create(created)
        return Response({"created": len(created)}, status=status.HTTP_201_CREATED)
    except (ValueError, TypeError, KeyError) as exc:
        return Response({"detail": str(exc)}, status=status.HTTP_400_BAD_REQUEST)


def _receipt_amount(lines):
    candidates = []
    for line in lines:
        match = re.search(r"(?:total|grand total|amount due|net total)\D{0,12}(\d+[,.]?\d{0,2})", line, re.I)
        if match:
            candidates.append(match.group(1).replace(",", ""))
    if candidates:
        return candidates[-1]
    values = re.findall(r"(?<!\d)(\d{1,7}[.]\d{2})(?!\d)", " ".join(lines))
    return values[-1] if values else ""


def _receipt_date(lines):
    for line in lines:
        match = re.search(r"(\d{4})[-/]([01]?\d)[-/]([0-3]?\d)|(\d{1,2})[-/]([01]?\d)[-/](\d{2,4})", line)
        if not match:
            continue
        groups = match.groups()
        if groups[0]:
            return f"{int(groups[0]):04d}-{int(groups[1]):02d}-{int(groups[2]):02d}"
        year = int(groups[5])
        year += 2000 if year < 100 else 0
        return f"{year:04d}-{int(groups[4]):02d}-{int(groups[3]):02d}"
    return str(timezone.localdate())


def _receipt_category(text):
    lowered = text.lower()
    for keyword, category in (
        ("rent", "rent"), ("utility", "utilities"), ("electric", "utilities"),
        ("transport", "transport"), ("fuel", "transport"), ("market", "supplies"),
        ("supply", "supplies"), ("salary", "salary"), ("wage", "salary"),
        ("advert", "marketing"),
    ):
        if keyword in lowered:
            return category
    return "other"


@api_view(["POST"])
@parser_classes([MultiPartParser])
@permission_classes([IsManagerOrModerator, HasActiveSubscription])
def expense_receipt_extract(request):
    upload = request.FILES.get("file")
    if not upload:
        return Response({"detail": "Attach a receipt image in the file field."}, status=status.HTTP_400_BAD_REQUEST)
    if not os.environ.get("AWS_ACCESS_KEY_ID") or not os.environ.get("AWS_SECRET_ACCESS_KEY"):
        return Response({"detail": "Receipt OCR is not configured. Set AWS Textract credentials on the backend."}, status=status.HTTP_503_SERVICE_UNAVAILABLE)
    try:
        import boto3

        client = boto3.client("textract", region_name=os.environ.get("AWS_REGION", "us-east-1"))
        result = client.detect_document_text(Document={"Bytes": upload.read()})
        lines = [block["Text"] for block in result.get("Blocks", []) if block.get("BlockType") == "LINE"]
        if not lines:
            return Response({"detail": "No text was detected in the receipt."}, status=status.HTTP_422_UNPROCESSABLE_ENTITY)
        text = " ".join(lines)
        return Response({
            "amount": _receipt_amount(lines),
            "date": _receipt_date(lines),
            "category": _receipt_category(text),
            "vendor_name": lines[0][:160] if lines else "",
            "description": next((line[:255] for line in lines if len(line) > 3), "Receipt expense"),
            "text": "\n".join(lines),
        })
    except Exception as exc:
        return Response({"detail": f"Receipt OCR failed: {exc}"}, status=status.HTTP_502_BAD_GATEWAY)
